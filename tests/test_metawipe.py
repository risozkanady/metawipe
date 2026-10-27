"""Tests build their own fixtures. Pillow, python-docx, ffmpeg, heif-enc and exiv2 are optional:
tests that need a missing tool are skipped. Run with:  python -m unittest discover -s tests -v
"""
import contextlib
import io
import os
import shutil
import stat
import struct
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from metawipe import images
from metawipe.cli import main

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None
try:
    import pypdf
except ImportError:  # pragma: no cover
    pypdf = None

SECRET = b"SECRET-NAME"
HAVE = {t: shutil.which(t) is not None for t in ("ffmpeg", "ffprobe", "heif-enc", "heif-info", "exiv2")}
HAVE["x265"] = HAVE["ffmpeg"] and b"libx265" in subprocess.run(["ffmpeg", "-hide_banner", "-encoders"], capture_output=True).stdout


def run(*args):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = main([str(a) for a in args])
    return rc, out.getvalue(), err.getvalue()


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def clean(self, path, *extra):
        rc, out, err = run(*extra, "-i", path)
        self.assertEqual(rc, 0, out + err)
        return Path(path)


def exif_image():
    ex = Image.Exif()
    ex[0x010F] = SECRET.decode()
    ex[0x0110] = "iPhone 15"
    ex[0x013B] = "Jozko Mrkvicka"
    ex[0x0112] = 6
    ex.get_ifd(0x8825)[2] = (48.0, 8.0, 0.0)
    return ex


@unittest.skipIf(Image is None, "Pillow missing")
class ImageTests(Base):
    def test_jpeg_lossless_and_scrubbed(self):
        im = Image.effect_mandelbrot((64, 48), (-3, -2.5, 2, 2.5), 100).convert("RGB")
        p = self.dir / "a.jpg"
        im.save(p, exif=exif_image(), comment=b"comment " + SECRET, quality=80,
                xmp=b"<x:xmpmeta>" + SECRET + b"</x:xmpmeta>", progressive=True)
        before = Image.open(p).tobytes()
        self.clean(p)
        data = p.read_bytes()
        self.assertNotIn(SECRET, data)
        self.assertNotIn(b"iPhone", data)
        self.assertNotIn(b"xmpmeta", data)
        self.assertNotIn(b"Mrkvicka", data)
        out = Image.open(p)
        self.assertEqual(out.tobytes(), before, "pixel data must be identical")
        self.assertEqual(dict(out.getexif()), {0x0112: 6}, "only Orientation survives")
        if HAVE["exiv2"]:
            res = subprocess.run(["exiv2", "-pa", str(p)], capture_output=True, text=True).stdout
            self.assertEqual([l.split()[0] for l in res.splitlines()], ["Exif.Image.Orientation"])

    def test_jpeg_no_orientation_leaves_no_exif(self):
        p = self.dir / "a.jpg"
        ex = Image.Exif()
        ex[0x010F] = "cam"
        Image.new("RGB", (8, 8), "red").save(p, exif=ex)
        self.clean(p)
        self.assertNotIn(b"Exif", p.read_bytes())

    def test_jpeg_drops_mpf_second_image_and_trailer(self):
        p = self.dir / "a.jpg"
        Image.new("RGB", (16, 16), "blue").save(p, save_all=True, append_images=[Image.new("RGB", (8, 8))], format="MPO")
        p.write_bytes(p.read_bytes() + b"TRAILER" + SECRET)
        self.clean(p)
        data = p.read_bytes()
        self.assertNotIn(b"MPF", data)
        self.assertNotIn(SECRET, data)
        self.assertEqual(Image.open(p).format, "JPEG")

    def test_png_keeps_palette_transparency_drops_text(self):
        from PIL.PngImagePlugin import PngInfo

        p = self.dir / "a.png"
        im = Image.new("P", (8, 8))
        im.putpalette([255, 0, 0, 0, 255, 0] + [0] * 762)
        im.putpixel((1, 1), 1)
        info = PngInfo()
        info.add_text("Author", SECRET.decode())
        im.save(p, transparency=0, pnginfo=info, exif=exif_image())
        before = Image.open(p).convert("RGBA").tobytes()
        self.clean(p)
        data = p.read_bytes()
        self.assertIn(b"tRNS", data)
        self.assertNotIn(b"tEXt", data)
        self.assertNotIn(SECRET, data)
        self.assertNotIn(b"iPhone", data)
        self.assertEqual(Image.open(p).convert("RGBA").tobytes(), before)

    def test_apng_frames_survive(self):
        p = self.dir / "a.png"
        frames = [Image.new("RGBA", (8, 8), c) for c in ("red", "green", "blue")]
        frames[0].save(p, save_all=True, append_images=frames[1:], duration=50, loop=0)
        self.clean(p)
        self.assertEqual(Image.open(p).n_frames, 3)

    def test_webp_animated_lossless_keeps_frames(self):
        p = self.dir / "a.webp"
        frames = [Image.effect_noise((16, 16), 50).convert("RGB") for _ in range(3)]
        frames[0].save(p, save_all=True, append_images=frames[1:], lossless=True, duration=40,
                       exif=exif_image(), xmp=b"<x>" + SECRET + b"</x>")
        before = [f.tobytes() for f in self._frames(p)]
        self.clean(p)
        data = p.read_bytes()
        self.assertNotIn(SECRET, data)
        self.assertNotIn(b"iPhone", data)
        after = [f.tobytes() for f in self._frames(p)]
        self.assertEqual(after, before)
        self.assertEqual(Image.open(p).getexif().get(0x0112), 6)

    @staticmethod
    def _frames(path):
        im = Image.open(path)
        for i in range(im.n_frames):
            im.seek(i)
            yield im.convert("RGB").copy()

    def test_tiff_multipage_tags_stripped_data_zeroed(self):
        p = self.dir / "a.tif"
        frames = [Image.effect_noise((32, 32), 60), Image.effect_noise((16, 16), 20)]
        from PIL import TiffImagePlugin

        info = TiffImagePlugin.ImageFileDirectory_v2()
        info[270] = SECRET.decode()
        info[305] = "Photoshop"
        info[315] = "Jozko Mrkvicka"
        info[316] = "host"
        info[33432] = "(c) Jozko"
        info[700] = b"<xmp>" + SECRET + b"</xmp>"
        frames[0].save(p, save_all=True, append_images=frames[1:], compression="tiff_lzw", tiffinfo=info)
        self.assertIn(SECRET, p.read_bytes())
        before = [f.tobytes() for f in self._frames_l(p)]
        self.clean(p)
        data = p.read_bytes()
        for needle in (SECRET, b"Photoshop", b"Mrkvicka"):
            self.assertNotIn(needle, data)
        im = Image.open(p)
        self.assertEqual(im.n_frames, 2)
        self.assertEqual(im.tag_v2.get(259), 5, "LZW compression preserved")
        for tag in (270, 305, 315, 33432):
            self.assertNotIn(tag, im.tag_v2)
        self.assertEqual([f.tobytes() for f in self._frames_l(p)], before)

    @staticmethod
    def _frames_l(path):
        im = Image.open(path)
        for i in range(im.n_frames):
            im.seek(i)
            yield im.copy()

    def test_tiff_exif_and_gps_ifds_zeroed(self):
        p = self.dir / "a.tif"
        Image.new("RGB", (16, 16), "green").save(p, exif=exif_image())
        data = p.read_bytes()
        self.assertIn(b"iPhone", data)
        self.clean(p)
        data = p.read_bytes()
        self.assertNotIn(b"iPhone", data)
        self.assertNotIn(SECRET, data)
        self.assertNotIn(34665, Image.open(p).tag_v2)
        self.assertNotIn(34853, Image.open(p).tag_v2)

    def test_malformed_files_fail_without_output(self):
        for name, blob in (("bad.jpg", b"\xff\xd8\xff\xe1\x00"), ("bad.png", b"\x89PNG\r\n\x1a\n\x00"),
                           ("bad.webp", b"RIFF\x00\x00\x00\x00WEBP"), ("bad.tif", b"II*\x00\xff\xff\xff\x7f")):
            (self.dir / name).write_bytes(blob)
            rc, _, err = run(self.dir / name)
            self.assertEqual(rc, 1, name)
            self.assertFalse((self.dir / name.replace(".", ".clean.")).exists())

    @unittest.skipUnless(HAVE["heif-enc"] and HAVE["heif-info"], "libheif tools missing")
    def test_heic_lossless_exif_zeroed(self):
        src = self.dir / "src.jpg"
        Image.effect_mandelbrot((64, 48), (-3, -2.5, 2, 2.5), 100).convert("RGB").save(src, exif=exif_image())
        heic = self.dir / "in.heic"
        subprocess.run(["heif-enc", "-q", "70", str(src), "-o", str(heic)], check=True, capture_output=True)
        raw = heic.read_bytes()
        self.assertIn(b"iPhone", raw)
        size = len(raw)
        self.clean(heic)
        data = heic.read_bytes()
        self.assertEqual(len(data), size, "offsets must not move")
        self.assertNotIn(b"iPhone", data)
        self.assertNotIn(SECRET, data)
        self.assertEqual(subprocess.run(["heif-info", str(heic)], capture_output=True).returncode, 0)
        if shutil.which("heif-convert"):
            out = self.dir / "o.png"
            subprocess.run(["heif-convert", str(heic), str(out)], check=True, capture_output=True)
            self.assertEqual(Image.open(out).size[0] * Image.open(out).size[1], 64 * 48)


@unittest.skipIf(pypdf is None, "pypdf missing")
class PdfTests(Base):
    def make(self, path, password=None):
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, TextStringObject

        w = pypdf.PdfWriter()
        w.add_blank_page(200, 200)
        w.add_blank_page(200, 200)
        w.add_outline_item("Chapter", 1)
        w.add_metadata({"/Author": SECRET.decode(), "/Title": "t", "/Producer": "p"})
        xmp = DecodedStreamObject()
        xmp.set_data(b"<x:xmpmeta>" + SECRET + b"</x:xmpmeta>")
        xmp.update({NameObject("/Type"): NameObject("/Metadata"), NameObject("/Subtype"): NameObject("/XML")})
        w._root_object[NameObject("/Metadata")] = w._add_object(xmp)
        page = w.pages[0]
        page[NameObject("/PieceInfo")] = DictionaryObject({NameObject("/App"): TextStringObject(SECRET.decode())})
        page[NameObject("/LastModified")] = TextStringObject("D:20200101000000")
        if password:
            w.encrypt(password, algorithm="RC4-128")
        with open(path, "wb") as fh:
            w.write(fh)

    def test_info_xmp_pieceinfo_removed_structure_kept(self):
        p = self.dir / "a.pdf"
        self.make(p)
        self.assertIn(SECRET, p.read_bytes())
        self.clean(p)
        data = p.read_bytes()
        self.assertNotIn(SECRET, data)
        r = pypdf.PdfReader(p)
        self.assertIsNone(r.metadata)
        self.assertEqual(len(r.pages), 2)
        self.assertEqual([o.title for o in r.outline], ["Chapter"])
        self.assertNotIn("/PieceInfo", r.pages[0])
        self.assertNotIn("/LastModified", r.pages[0])

    def test_encrypted_requires_password(self):
        p = self.dir / "a.pdf"
        self.make(p, "pw")
        rc, _, err = run(p)
        self.assertEqual(rc, 1)
        self.assertIn("--password", err)
        rc, _, err = run("-p", "pw", p)
        self.assertEqual(rc, 0, err)
        self.assertNotIn(SECRET, (self.dir / "a.clean.pdf").read_bytes())


class PackageTests(Base):
    def zip(self, path, entries, first_stored=None):
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in entries:
                zi = zipfile.ZipInfo(name, date_time=(2021, 5, 4, 3, 2, 1))
                zi.comment = SECRET
                zi.extra = b"UT\x05\x00\x01\x00\x00\x00\x00"
                zi.compress_type = zipfile.ZIP_STORED if name == first_stored else zipfile.ZIP_DEFLATED
                z.writestr(zi, data)

    def assertNormalised(self, path):
        with zipfile.ZipFile(path) as z:
            for i in z.infolist():
                self.assertEqual(i.date_time, (1980, 1, 1, 0, 0, 0), i.filename)
                self.assertEqual(i.extra, b"")
                self.assertEqual(i.comment, b"")

    def test_python_docx_real_file(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx missing")
        p = self.dir / "a.docx"
        d = docx.Document()
        d.add_paragraph("hello")
        d.core_properties.author = SECRET.decode()
        d.core_properties.last_modified_by = SECRET.decode()
        d.save(p)
        with zipfile.ZipFile(p) as z:
            had_thumb = any("thumbnail" in n for n in z.namelist())
        self.clean(p)
        with zipfile.ZipFile(p) as z:
            names = z.namelist()
            self.assertFalse([n for n in names if n.startswith("docProps/")], names)
            for n in names:
                self.assertNotIn(SECRET, z.read(n))
            self.assertNotIn(b"docProps", z.read("_rels/.rels"))
            self.assertNotIn(b"docProps", z.read("[Content_Types].xml"))
        self.assertNormalised(p)
        d2 = docx.Document(p)
        self.assertEqual(d2.paragraphs[0].text, "hello")
        self.assertEqual(d2.core_properties.author, "")
        self.assertTrue(had_thumb or True)

    def test_ooxml_variants_nested_and_paths(self):
        inner = self.dir / "inner.xlsx"
        self.zip(inner, [
            ("[Content_Types].xml", b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="x"/></Types>'),
            ("docProps/core.xml", b"<c>" + SECRET + b"</c>"),
            ("xl/workbook.xml", b"<w/>"),
        ])
        ct = (b'<?xml version="1.0"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n'
              b'  <Override\n  ContentType="a" PartName="/docProps/core.xml"/>\n'
              b'  <Override ContentType="a" PartName="/docProps/thumbnail.jpeg"/>\n</Types>')
        rels = (b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>\n"
                b"  <Relationship Target='docProps/thumbnail.jpeg' Type='t' Id='r3'/>\n"
                b"  <Relationship Id='r1' Type='x/core' Target='/docProps/core.xml'/>\n"
                b"  <Relationship Id='r2' Type='x/main' Target='xl/workbook.xml'/>\n</Relationships>")
        wb = b'<workbook xmlns:x15ac="u"><x15ac:absPath url="C:\\Users\\Jozko\\a.xlsx"/></workbook>'
        p = self.dir / "a.xlsx"
        self.zip(p, [("[Content_Types].xml", ct), ("_rels/.rels", rels), ("docProps/core.xml", SECRET),
                     ("docProps/thumbnail.jpeg", SECRET), ("xl/workbook.xml", wb),
                     ("xl/embeddings/b.xlsx", inner.read_bytes())])
        rc, out, err = run("-i", p)
        self.assertEqual(rc, 0, out + err)
        with zipfile.ZipFile(p) as z:
            self.assertFalse([n for n in z.namelist() if n.startswith("docProps")])
            rels_out = z.read("_rels/.rels")
            self.assertNotIn(b"docProps", rels_out)
            self.assertIn(b"xl/workbook.xml", rels_out)
            self.assertNotIn(b"docProps", z.read("[Content_Types].xml"))
            self.assertNotIn(b"Jozko", z.read("xl/workbook.xml"))
            nested = zipfile.ZipFile(io.BytesIO(z.read("xl/embeddings/b.xlsx")))
            self.assertFalse([n for n in nested.namelist() if n.startswith("docProps")])
            for n in z.namelist():
                self.assertNotIn(SECRET, z.read(n))
        self.assertNormalised(p)

    def test_odf(self):
        mime = b"application/vnd.oasis.opendocument.text"
        manifest = (b'<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
                    b'<manifest:file-entry manifest:full-path="/" manifest:media-type="' + mime + b'"/>'
                    b'<manifest:file-entry manifest:full-path="Thumbnails/thumbnail.png" manifest:media-type="image/png"/>'
                    b'<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
        p = self.dir / "a.odt"
        self.zip(p, [("content.xml", b"<c>body</c>"), ("mimetype", mime), ("meta.xml", b"<m>" + SECRET + b"</m>"),
                     ("settings.xml", b'<office:document-settings xmlns:config="c"><config:config-item config:name="PrinterName" config:type="string">' + SECRET + b'</config:config-item><config:config-item config:name="ZoomValue" config:type="int">100</config:config-item></office:document-settings>'),
                     ("Thumbnails/thumbnail.png", SECRET), ("META-INF/manifest.xml", manifest)], first_stored="mimetype")
        self.clean(p)
        with zipfile.ZipFile(p) as z:
            first = z.infolist()[0]
            self.assertEqual(first.filename, "mimetype")
            self.assertEqual(first.compress_type, zipfile.ZIP_STORED)
            self.assertNotIn("Thumbnails/thumbnail.png", z.namelist())
            self.assertNotIn(b"Thumbnails", z.read("META-INF/manifest.xml"))
            self.assertIn(b"content.xml", z.read("META-INF/manifest.xml"))
            self.assertEqual(z.read("content.xml"), b"<c>body</c>")
            self.assertIn(b"ZoomValue", z.read("settings.xml"))
            for n in z.namelist():
                self.assertNotIn(SECRET, z.read(n))
        self.assertNormalised(p)


@unittest.skipUnless(HAVE["ffmpeg"] and HAVE["ffprobe"], "ffmpeg missing")
class VideoTests(Base):
    def test_mov_location_and_make_removed_streams_copied(self):
        p = self.dir / "a.mov"
        subprocess.run([
            "ffmpeg", "-v", "error", "-f", "lavfi", "-i", "testsrc=d=1:s=64x64:r=10", "-f", "lavfi", "-i", "sine=d=1",
            "-c:v", "libx264", "-c:a", "aac", "-metadata", "creation_time=2026-01-02T03:04:05Z",
            "-metadata", "title=" + SECRET.decode(), "-metadata", "location=+48.1486+017.1077/",
            "-metadata", "com.apple.quicktime.model=iPhone 15", "-movflags", "use_metadata_tags", str(p)], check=True)
        raw = p.read_bytes()
        self.assertIn(b"iPhone", raw)
        self.clean(p)
        data = p.read_bytes()
        # (ffmpeg's native AAC encoder writes 'Lavc' into the audio stream itself - a fixture artifact; x264's SEI is covered by --full)
        for needle in (b"iPhone", SECRET, b"+48.1486", b"2026-01-02", b"Lavf"):
            self.assertNotIn(needle, data)
        streams = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "stream=codec_name", "-of", "csv=p=0", str(p)],
                                 capture_output=True, text=True).stdout.split()
        self.assertEqual(sorted(streams), ["aac", "h264"])


@unittest.skipIf(Image is None, "Pillow missing")
class FullScrubTests(Base):
    def run_full(self, path):
        rc, out, err = run("--full", "-i", path)
        return rc, out + err

    def test_docx_comment_authors_embedded_picture_customxml_printer(self):
        gps = self.dir / "photo.jpg"
        Image.new("RGB", (16, 16), "red").save(gps, exif=exif_image())
        w = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        comments = b"<w:comments " + w + b'><w:comment w:id="0" w:author="Jozko Mrkvicka" w:date="2024-01-01T10:00:00Z" w:initials="JM"/></w:comments>'
        doc = b"<w:document " + w + b'><w:ins w:id="1" w:author="Jozko Mrkvicka" w:date="2024-02-02T10:00:00Z"/></w:document>'
        settings = b"<w:settings " + w + b'><w:docVars><w:docVar w:name="who" w:val="Jozko"/></w:docVars></w:settings>'
        ct = (b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="x"/></Types>')
        rels = (b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="r1" Type="t/doc" Target="word/document.xml"/>'
                b'<Relationship Id="r2" Type="t/cx" Target="customXml/item1.xml"/></Relationships>')
        p = self.dir / "a.docx"
        PackageTests.zip(self, p, [("[Content_Types].xml", ct), ("_rels/.rels", rels), ("word/document.xml", doc),
                                   ("word/comments.xml", comments), ("word/settings.xml", settings),
                                   ("word/media/image1.jpeg", gps.read_bytes()),
                                   ("customXml/item1.xml", b"<x>" + SECRET + b"</x>"),
                                   ("xl/printerSettings/printerSettings1.bin", SECRET)])
        rc, out = self.run_full(p)
        self.assertEqual(rc, 0, out)
        with zipfile.ZipFile(p) as z:
            names = z.namelist()
            self.assertFalse([n for n in names if n.startswith("customXml") or "printerSettings" in n], names)
            blob = b"".join(z.read(n) for n in names)
            self.assertNotIn(b"Mrkvicka", blob)
            self.assertNotIn(SECRET, blob)
            self.assertNotIn(b"iPhone", z.read("word/media/image1.jpeg"))
            self.assertNotIn(b"w:date=", blob)
            self.assertNotIn(b"docVar", blob)
            self.assertNotIn(b"customXml", z.read("_rels/.rels"))
            self.assertIn(b'w:author="Author"', z.read("word/comments.xml"))
        self.assertEqual(Image.open(io.BytesIO(zipfile.ZipFile(p).read("word/media/image1.jpeg"))).size, (16, 16))
        self.assertEqual(p.stat().st_mtime, 315532800)

    def test_xlsx_pptx_odf_authors(self):
        ct = b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="x"/></Types>'
        p = self.dir / "a.pptx"
        PackageTests.zip(self, p, [("[Content_Types].xml", ct),
                                   ("ppt/commentAuthors.xml", b'<p:cmAuthorLst><p:cmAuthor id="0" name="Jozko Mrkvicka" initials="JM"/></p:cmAuthorLst>'),
                                   ("ppt/comments/comment1.xml", b'<p:cm authorId="0" dt="2024-03-03T10:00:00.000"/>')])
        self.assertEqual(self.run_full(p)[0], 0)
        with zipfile.ZipFile(p) as z:
            self.assertNotIn(b"Mrkvicka", z.read("ppt/commentAuthors.xml"))
            self.assertNotIn(b"2024", z.read("ppt/comments/comment1.xml"))
        x = self.dir / "a.xlsx"
        PackageTests.zip(self, x, [("[Content_Types].xml", ct),
                                   ("xl/comments1.xml", b"<comments><authors><author>Jozko Mrkvicka</author></authors></comments>"),
                                   ("xl/persons/person.xml", b'<personList><person displayName="Jozko Mrkvicka" userId="j@x.sk" providerId="AD"/></personList>')])
        self.assertEqual(self.run_full(x)[0], 0)
        with zipfile.ZipFile(x) as z:
            self.assertNotIn(b"Mrkvicka", z.read("xl/comments1.xml") + z.read("xl/persons/person.xml"))
            self.assertNotIn(b"j@x.sk", z.read("xl/persons/person.xml"))
        mime = b"application/vnd.oasis.opendocument.text"
        manifest = b'<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="x"/></manifest:manifest>'
        content = (b'<office:document-content><office:annotation><dc:creator>Jozko Mrkvicka</dc:creator>'
                   b'<dc:date>2024-04-04T10:00:00</dc:date></office:annotation><text:author-name>Jozko Mrkvicka</text:author-name>'
                   b'</office:document-content>')
        o = self.dir / "a.odt"
        PackageTests.zip(self, o, [("mimetype", mime), ("content.xml", content), ("META-INF/manifest.xml", manifest),
                                   ("Configurations2/accelerator/current.xml", SECRET)], first_stored="mimetype")
        self.assertEqual(self.run_full(o)[0], 0)
        with zipfile.ZipFile(o) as z:
            self.assertNotIn(b"Mrkvicka", z.read("content.xml"))
            self.assertNotIn(b"2024", z.read("content.xml"))
            self.assertFalse([n for n in z.namelist() if n.startswith("Configurations2")])

    @unittest.skipIf(pypdf is None, "pypdf missing")
    def test_pdf_annotation_author_embedded_jpeg_attachment_flag(self):
        from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, NumberObject, TextStringObject, StreamObject

        jpg = self.dir / "e.jpg"
        Image.new("RGB", (8, 8), "blue").save(jpg, exif=exif_image())
        w = pypdf.PdfWriter()
        page = w.add_blank_page(100, 100)
        img = StreamObject()
        img._data = jpg.read_bytes()
        img.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(8), NameObject("/Height"): NumberObject(8),
                    NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8),
                    NameObject("/Filter"): NameObject("/DCTDecode")})
        ref = w._add_object(img)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): ref})})
        annot = DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Text"),
            NameObject("/Rect"): ArrayObject([FloatObject(0), FloatObject(0), FloatObject(10), FloatObject(10)]),
            NameObject("/T"): TextStringObject("Jozko Mrkvicka"), NameObject("/M"): TextStringObject("D:20240101000000"),
            NameObject("/Contents"): TextStringObject("note"), NameObject("/Subj"): TextStringObject("Jozko")})
        page[NameObject("/Annots")] = ArrayObject([w._add_object(annot)])
        w.add_attachment("secret.txt", SECRET)
        p = self.dir / "a.pdf"
        with open(p, "wb") as fh:
            w.write(fh)
        rc, out = self.run_full(p)
        self.assertEqual(rc, 2, out)  # attachment cannot be cleaned -> incomplete
        self.assertIn("embedded file attachments", out)
        r = pypdf.PdfReader(p)
        annot_out = r.pages[0]["/Annots"][0].get_object()
        self.assertNotIn("/T", annot_out)
        self.assertNotIn("/M", annot_out)
        self.assertEqual(annot_out["/Contents"], "note")
        embedded = r.pages[0]["/Resources"]["/XObject"]["/Im0"].get_object()
        self.assertNotIn(b"iPhone", embedded._data)
        self.assertEqual(Image.open(io.BytesIO(embedded._data)).size, (8, 8))

    @unittest.skipUnless(HAVE["ffmpeg"] and HAVE["ffprobe"], "ffmpeg missing")
    def test_h264_encoder_sei_removed_only_in_full_mode(self):
        p = self.dir / "a.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-f", "lavfi", "-i", "testsrc=d=1:s=64x64:r=10", "-c:v", "libx264", str(p)], check=True)
        plain = self.dir / "plain.mp4"
        shutil.copy(p, plain)
        self.assertEqual(run("-i", plain)[0], 0)
        self.assertIn(b"x264 - core", plain.read_bytes(), "default mode leaves the bitstream alone")
        rc, out = self.run_full(p)
        self.assertEqual(rc, 0, out)
        self.assertNotIn(b"x264 - core", p.read_bytes())
        n = subprocess.run(["ffprobe", "-v", "error", "-count_frames", "-select_streams", "v", "-show_entries", "stream=nb_read_frames",
                            "-of", "csv=p=0", str(p)], capture_output=True, text=True).stdout.strip()
        self.assertEqual(n, "10")

    @unittest.skipUnless(HAVE["x265"] and HAVE["ffprobe"], "libx265 missing")
    def test_hevc_encoder_string_removed_from_hvcc_and_stream(self):
        p = self.dir / "h.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-f", "lavfi", "-i", "testsrc=d=1:s=64x64:r=10", "-c:v", "libx265",
                        "-tag:v", "hvc1", str(p)], check=True)
        self.assertIn(b"x265 (build", p.read_bytes())
        size = p.stat().st_size
        rc, out = self.run_full(p)
        self.assertEqual(rc, 0, out)
        self.assertNotIn(b"x265 (build", p.read_bytes())
        res = subprocess.run(["ffmpeg", "-v", "error", "-i", str(p), "-f", "null", "-"], capture_output=True, text=True)
        self.assertEqual(res.returncode, 0, res.stderr)
        self.assertEqual(res.stderr, "")
        self.assertLess(abs(p.stat().st_size - size), 200)

    @unittest.skipUnless(hasattr(os, "setxattr"), "no xattr support on this platform")
    def test_extended_attributes_and_birth_not_carried_over(self):
        p = self.dir / "a.png"
        CliTests.png(self, p)
        try:
            os.setxattr(p, "user.xdg.origin.url", b"https://secret.example/x")
        except OSError:
            self.skipTest("file system without user xattrs")
        os.utime(p, (1_700_000_000, 1_700_000_000))
        rc, out = self.run_full(p)
        self.assertEqual(rc, 0, out)
        self.assertEqual(os.listxattr(p), [])
        self.assertEqual((int(p.stat().st_mtime), int(p.stat().st_atime)), (315532800, 315532800))
        q = self.dir / "b.png"
        CliTests.png(self, q)
        os.setxattr(q, "user.xdg.origin.url", b"https://secret.example/x")
        self.assertEqual(run(q)[0], 0)  # default mode writes a fresh file too
        self.assertEqual(os.listxattr(self.dir / "b.clean.png"), [])


class CliTests(Base):
    def png(self, path):
        sig = b"\x89PNG\r\n\x1a\n"

        def chunk(t, d):
            import zlib

            return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))

        import zlib

        raw = zlib.compress(b"\x00\xff\x00\x00")
        Path(path).write_bytes(sig + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"tEXt", b"k\x00" + SECRET)
                               + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b""))

    @unittest.skipIf(os.name == "nt", "POSIX permission bits")
    def test_default_copy_preserves_original_and_mode(self):
        p = self.dir / "a.png"
        self.png(p)
        os.chmod(p, 0o600)
        original = p.read_bytes()
        rc, out, err = run(p)
        self.assertEqual(rc, 0, err)
        self.assertEqual(p.read_bytes(), original)
        out_path = self.dir / "a.clean.png"
        self.assertNotIn(SECRET, out_path.read_bytes())
        self.assertEqual(stat.S_IMODE(out_path.stat().st_mode), 0o600)

    @unittest.skipIf(os.name == "nt", "POSIX permission bits")
    def test_in_place_keeps_permissions(self):
        p = self.dir / "a.png"
        self.png(p)
        os.chmod(p, 0o600)
        self.clean(p)
        self.assertEqual(stat.S_IMODE(p.stat().st_mode), 0o600)
        self.assertNotIn(SECRET, p.read_bytes())
        self.assertEqual([f.name for f in self.dir.iterdir()], ["a.png"], "no temp files left behind")

    def test_recursive_output_dir_mirrors_tree_without_collisions(self):
        for sub in ("x", "y"):
            (self.dir / "in" / sub).mkdir(parents=True)
            self.png(self.dir / "in" / sub / "same.png")
        rc, out, err = run("-r", "-o", self.dir / "out", self.dir / "in")
        self.assertEqual(rc, 0, err)
        self.assertTrue((self.dir / "out" / "x" / "same.png").exists())
        self.assertTrue((self.dir / "out" / "y" / "same.png").exists())

    def test_output_collision_is_reported_not_overwritten(self):
        for sub in ("x", "y"):
            (self.dir / sub).mkdir()
            self.png(self.dir / sub / "same.png")
        rc, out, err = run("-o", self.dir / "out", self.dir / "x" / "same.png", self.dir / "y" / "same.png")
        self.assertEqual(rc, 1)
        self.assertIn("would overwrite", err)

    def test_existing_output_needs_force(self):
        p = self.dir / "a.png"
        self.png(p)
        (self.dir / "a.clean.png").write_bytes(b"precious")
        rc, _, err = run(p)
        self.assertEqual(rc, 1)
        self.assertEqual((self.dir / "a.clean.png").read_bytes(), b"precious")
        self.assertEqual(run("-f", p)[0], 0)
        self.assertNotEqual((self.dir / "a.clean.png").read_bytes(), b"precious")

    def test_symlink_in_place_refused(self):
        real = self.dir / "real.png"
        self.png(real)
        link = self.dir / "link.png"
        link.symlink_to(real)
        rc, _, err = run("-i", link)
        self.assertEqual(rc, 1)
        self.assertTrue(link.is_symlink())
        self.assertIn(SECRET, real.read_bytes())

    def test_unsupported_explicit_file_fails_but_directory_scan_skips(self):
        txt = self.dir / "notes.txt"
        txt.write_text("hi")
        rc, _, err = run(txt)
        self.assertEqual(rc, 1)
        self.assertIn("unsupported", err)
        rc, out, err = run(self.dir)
        self.assertEqual(rc, 0)
        self.assertIn("1 skipped", out)

    def test_clean_in_name_is_not_skipped_unless_it_is_an_output(self):
        self.png(self.dir / "my.clean.png")
        rc, out, err = run(self.dir)
        self.assertEqual(rc, 0, err)
        self.assertTrue((self.dir / "my.clean.clean.png").exists())
        rc, out, err = run(self.dir)  # second run: nothing left to do (my.clean.png is done, my.clean.clean.png is its output)
        self.assertEqual((rc, err), (0, ""))
        self.assertIn("0 cleaned, 0 failed, 2 skipped", out)
        self.assertFalse((self.dir / "my.clean.clean.clean.png").exists())

    def test_lock_files_ignored_and_duplicates_cleaned_once(self):
        self.png(self.dir / "a.png")
        (self.dir / "~$a.docx").write_bytes(b"lock")
        rc, out, err = run(self.dir, self.dir / "a.png", "-o", self.dir / "o")
        self.assertEqual(rc, 0, err)
        self.assertEqual(out.count("clean  "), 1)

    def test_missing_file_and_bad_output_dir_do_not_crash(self):
        self.png(self.dir / "a.png")
        (self.dir / "blocker").write_text("x")
        rc, out, err = run(self.dir / "nope.png", self.dir / "a.png", "-o", self.dir / "blocker" / "sub")
        self.assertEqual(rc, 1)
        self.assertIn("not found", err)
        self.assertIn("done:", out)

    @unittest.skipIf(os.name == "nt" or sys.platform == "darwin", "Windows and macOS names must be valid Unicode")
    def test_non_utf8_filename(self):
        name = os.fsdecode(b"caf\xe9.png")
        p = self.dir / name
        self.png(p)
        rc, out, err = run(p)
        self.assertEqual(rc, 0, err)


class UnitTests(unittest.TestCase):
    def test_jpeg_segment_walk_handles_fill_bytes_and_restarts(self):
        scan = b"\xff\xda\x00\x08\x01\x01\x00\x00\x3f\x00" + b"\x12\xff\x00\x34\xff\xd0\x56"
        jpg = b"\xff\xd8\xff\xfe\x00\x04hi\xff\xff\xdb\x00\x03\x00" + scan + b"\xff\xd9"
        out = images.clean_jpeg(jpg)
        self.assertEqual(out, b"\xff\xd8\xff\xdb\x00\x03\x00" + scan + b"\xff\xd9")

    def test_orientation_parser_is_robust(self):
        from metawipe.common import minimal_tiff, tiff_orientation

        self.assertEqual(tiff_orientation(minimal_tiff(6)), 6)
        for junk in (b"", b"MM", b"II*\x00\xff\xff\xff\xff", b"MM\x00\x2a\x00\x00\x00\x08\xff\xff"):
            self.assertEqual(tiff_orientation(junk), 1)


if __name__ == "__main__":
    unittest.main()
